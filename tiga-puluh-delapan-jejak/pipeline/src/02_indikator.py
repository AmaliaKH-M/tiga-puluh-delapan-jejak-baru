"""STEP 1c — Indikator sosial-ekonomi kab/kota dan provinsi (format long).

Output (data/processed/):
  indikator_kabkota_long.csv  : 38 kab/kota (tanpa baris JAWA TIMUR)
  indikator_provinsi_long.csv : Provinsi Jawa Timur (kode 3500)
  inflasi_jatim_2025_long.csv : D19, inflasi y-on-y per kelompok/subkelompok

Kolom indikator: kode_wilayah, nama_wilayah, tahun, indikator, nilai, satuan,
sumber_id, flag. Nilai NA tidak dibuang; diberi flag (mis. 'tidak_tersedia').

Aturan konversi (lihat reports/validation_report.md):
  - String '1.234,56' -> 1234.56
  - Sheet PPK & penduduk: angka tersimpan sebagai float dengan titik ribuan
    yang hilang (6.775 = 6.775 ribu; 586.11 = 586.110). Dikali 1000 hanya jika
    nilai < batas yang tidak ambigu (diuji di 05_validate.py).
  - Kepadatan kolom 2019: 16 nilai terbaca 1.135 dst. -> x1000 jika < 20.
  - Kolom R sheet kepadatan (tanpa header, bergeser 1 baris) diabaikan.
"""
import openpyxl
import pandas as pd
from common import (XLSX, UHH_LF, CONFIG, PROCESSED, YEARS, id_num,
                    ensure_dirs, load_kode)

# indikator: (sheet, sumber_id, satuan, batas_float_x1000)
SPEC = {
    "persen_penduduk_miskin": ("%miskin kab kot", "D03", "persen", None),
    "jumlah_penduduk_miskin": ("jml pen miskin (rb jiwa)", "D04", "ribu jiwa", None),
    "pengeluaran_per_kapita_disesuaikan": ("per kapita", "D05", "ribu rupiah/orang/tahun", 100),
    "ipm": ("ipm", "D06", "indeks", None),
    "rls": ("rls", "D07", "tahun", None),
    "hls": ("hls", "D08", "tahun", None),
    "uhh": ("uhh", "D09", "tahun", None),
    "tpt": ("tpt", "D10", "persen", None),
    "pdrb_per_kapita_adhb": ("adhb perkapita", "D11", "ribu rupiah", None),
    "laju_pertumbuhan_pdrb_adhk": ("laju adhk", "D12", "persen", None),
    "jumlah_penduduk": ("jml penduduk", "D13", "jiwa", 1000),
    "kepadatan_penduduk": ("kepadatan penduduk", "D14", "jiwa/km2", 20),
    "gini_ratio": ("gini ratio", "D15", "indeks", None),
}
PROV_NAMES = {"JAWA TIMUR", "Jawa Timur"}


def read_wide(wb, sheet, limit):
    ws = wb[sheet]
    years = [int(ws.cell(4, c).value) for c in range(2, 18)]
    assert years == YEARS, f"header tahun {sheet} tidak sesuai: {years}"
    out = {}
    for r in range(5, 44):  # 38 kab/kota + baris provinsi; baris lain (catatan, sampah) diabaikan
        nama = str(ws.cell(r, 1).value).strip()
        vals = []
        for c in range(2, 18):
            v = ws.cell(r, c).value
            if isinstance(v, float) and limit is not None and v < limit:
                v = round(v * 1000, 3)
            else:
                v = id_num(v)
            vals.append(v)
        out[nama] = vals
    return out


def add_flag(cur, new):
    return new if not cur else f"{cur};{new}"


def main():
    ensure_dirs()
    kode = load_kode()
    name2kode = dict(zip(kode.nama_tabel, kode.kode_wilayah))
    kode2nama = dict(zip(kode.kode_wilayah, kode.nama_wilayah))
    corr = pd.read_csv(CONFIG / "corrections.csv", dtype=str)
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    uhh_lf = pd.read_csv(UHH_LF, index_col=0)
    uhh_lf.columns = uhh_lf.columns.astype(int)

    rows = []
    for ind, (sheet, sid, sat, lim) in SPEC.items():
        data = read_wide(wb, sheet, lim)
        assert len(data) == 39, f"{sheet}: {len(data)} baris"
        for nama, vals in data.items():
            is_prov = nama in PROV_NAMES
            k = "3500" if is_prov else name2kode[nama]
            for th, v in zip(YEARS, vals):
                flag = ""
                if ind == "laju_pertumbuhan_pdrb_adhk" and th == 2010:
                    continue  # tahun dasar: laju tidak terdefinisi
                if ind == "uhh" and th >= 2020:
                    key = "JAWA TIMUR" if is_prov else nama
                    v = float(uhh_lf.loc[key, th])
                    flag = "uhh_lf_sp2020"
                if ind in ("uhh", "ipm") and th <= 2019:
                    flag = add_flag(flag, "patahan_seri_metode_uhh")
                if v is None:
                    flag = add_flag(flag, "tidak_tersedia")
                if ind == "jumlah_penduduk" and th in (2020, 2024):
                    flag = add_flag(flag, "patahan_seri")
                if ind == "kepadatan_penduduk" and th <= 2021:
                    flag = add_flag(flag, "luas_acuan_berbeda")
                if (ind in ("persen_penduduk_miskin", "jumlah_penduduk_miskin")
                        and nama == "Kota Probolinggo" and th in (2010, 2011)):
                    flag = add_flag(flag, "perlu_verifikasi")
                # koreksi nilai terdokumentasi
                c = corr[(corr.jenis == "koreksi_nilai") & (corr.indikator == ind) &
                         (corr.nama_wilayah_tabel == nama) & (corr.tahun == str(th))]
                if len(c):
                    assert abs(v - float(c.nilai_lama.iloc[0])) < 0.5, "nilai_lama tidak cocok"
                    v = float(c.nilai_baru.iloc[0])
                    flag = add_flag(flag, f"dikoreksi_{c.id.iloc[0]}")
                rows.append((k, ind, th, v, sat, sid, flag))

    # D18 makanan / bukan makanan (2020-2025, rupiah/kapita/bulan)
    ws = wb["makanan non mkn"]
    yrs18 = [int(ws.cell(4, c).value) for c in range(2, 8)]
    blocks = {"pengeluaran_makanan": 5, "pengeluaran_bukan_makanan": 47}
    d18 = {}
    for ind, r0 in blocks.items():
        for r in range(r0, r0 + 39):
            nama = str(ws.cell(r, 1).value).strip()
            k = "3500" if nama in PROV_NAMES else name2kode[nama]
            for j, th in enumerate(yrs18):
                v = id_num(ws.cell(r, 2 + j).value)
                d18[(k, ind, th)] = v
    for (k, ind, th), v in d18.items():
        flag = "perlu_verifikasi" if (k == "3579" and ind == "pengeluaran_bukan_makanan"
                                      and th in (2023, 2025)) else ""
        rows.append((k, ind, th, v, "rupiah/kapita/bulan", "D18", flag))
    for k in {key[0] for key in d18}:
        for th in yrs18:
            a, b = d18[(k, "pengeluaran_makanan", th)], d18[(k, "pengeluaran_bukan_makanan", th)]
            flag = "turunan_makanan_plus_bukan_makanan"
            if k == "3579" and th in (2023, 2025):
                flag += ";perlu_verifikasi"
            rows.append((k, "pengeluaran_total", th, round(a + b, 2),
                         "rupiah/kapita/bulan", "D18", flag))

    df = pd.DataFrame(rows, columns=["kode_wilayah", "indikator", "tahun", "nilai",
                                     "satuan", "sumber_id", "flag"])
    df["nama_wilayah"] = df.kode_wilayah.map(kode2nama).fillna("Provinsi Jawa Timur")
    cols = ["kode_wilayah", "nama_wilayah", "tahun", "indikator", "nilai", "satuan", "sumber_id", "flag"]
    df = df[cols]
    kab = df[df.kode_wilayah != "3500"].sort_values(["indikator", "kode_wilayah", "tahun"])
    prov = df[df.kode_wilayah == "3500"]

    # D01: seri provinsi % miskin 2005-2025 (menggantikan baris provinsi D03 untuk indikator ini)
    ws = wb["% pen miskin menurut tahun"]
    d01 = []
    for r in range(4, 25):
        th, v = ws.cell(r, 1).value, ws.cell(r, 2).value
        d01.append(("3500", "Provinsi Jawa Timur", int(th), "persen_penduduk_miskin",
                    id_num(v), "persen", "D01", "perlu_verifikasi" if th < 2010 else ""))
    d01 = pd.DataFrame(d01, columns=cols)
    prov_d03 = prov[prov.indikator == "persen_penduduk_miskin"].set_index("tahun").nilai
    chk = d01[d01.tahun >= 2010].set_index("tahun").nilai
    assert (abs(chk - prov_d03) < 1e-9).all(), "D01 != baris provinsi D03 untuk 2010-2025"
    prov = pd.concat([prov[prov.indikator != "persen_penduduk_miskin"], d01])
    prov = prov.sort_values(["indikator", "tahun"])

    kab.to_csv(PROCESSED / "indikator_kabkota_long.csv", index=False)
    prov.to_csv(PROCESSED / "indikator_provinsi_long.csv", index=False)

    # D19 inflasi 2025 (kode dibangun ulang sebagai teks: nol di depan hilang di Excel)
    ws = wb["inflasi tahunan "]
    bulan = [ws.cell(3, c).value for c in range(3, 15)]
    inf, kel = [], None
    for r in range(4, ws.max_row + 1):
        kv, nama = ws.cell(r, 1).value, ws.cell(r, 2).value
        if kv is None or nama is None:
            continue
        kv = int(kv)
        if kv == 0:
            kode_txt, level, kel = "00", "umum", 0
        elif kel is not None and kv == kel + 1 and kv <= 11:
            kel = kv
            kode_txt, level = f"{kv:02d}", "kelompok"
        else:
            kode_txt, level = (f"0{kv}" if kel < 10 else str(kv)), "subkelompok"
        for j, b in enumerate(bulan):
            raw = ws.cell(r, 3 + j).value
            v = id_num(raw)
            flag = "mendekati_nol" if str(raw).strip() in ("\\~0", "~0") else ""
            inf.append((kode_txt, level, str(nama).strip(), 2025, j + 1, b, v, flag))
    pd.DataFrame(inf, columns=["kode_kelompok", "level", "nama_kelompok", "tahun", "bulan",
                               "nama_bulan", "inflasi_yoy_persen", "flag"]
                 ).to_csv(PROCESSED / "inflasi_jatim_2025_long.csv", index=False)
    print(f"[02] indikator kab/kota: {len(kab)} baris | provinsi: {len(prov)} | inflasi: {len(inf)}")


if __name__ == "__main__":
    main()
