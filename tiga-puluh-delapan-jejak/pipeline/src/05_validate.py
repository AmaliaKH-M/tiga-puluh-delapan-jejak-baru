"""STEP 2 — Validasi seluruh data processed. Menulis reports/validation_report.md.
Keluar dengan kode 1 jika ada pemeriksaan FAIL (agar mudah dipakai di CI)."""
import json
import sys
import numpy as np
import pandas as pd
import openpyxl
from common import PROCESSED, GEO, INTERIM, REPORTS, XLSX, KODE_LU, load_kode

results = []


def check(group, name, ok, detail=""):
    results.append((group, name, "PASS" if ok else "FAIL", detail))


def warn(group, name, detail):
    results.append((group, name, "CATATAN", detail))


def main():
    kode = load_kode()
    ind = pd.read_csv(PROCESSED / "indikator_kabkota_long.csv", dtype={"kode_wilayah": str})
    prov = pd.read_csv(PROCESSED / "indikator_provinsi_long.csv", dtype={"kode_wilayah": str})
    pdrb = pd.read_csv(PROCESSED / "pdrb_lu_long.csv", dtype={"kode_wilayah": str})
    tot = pd.read_csv(PROCESSED / "pdrb_total_kabkota.csv", dtype={"kode_wilayah": str})
    expected = {f"{k}" for k in list(range(3501, 3530)) + list(range(3571, 3580))}

    # A. Kode wilayah
    g = "A. Kode wilayah"
    check(g, "tepat 38 kab/kota", len(kode) == 38, f"{len(kode)} baris")
    check(g, "string 4 digit", kode.kode_wilayah.str.fullmatch(r"\d{4}").all())
    check(g, "unik & tidak NULL", kode.kode_wilayah.is_unique and kode.kode_wilayah.notna().all())
    check(g, "himpunan = 3501–3529, 3571–3579", set(kode.kode_wilayah) == expected)

    # B. Indikator kab/kota
    g = "B. Indikator kab/kota"
    check(g, "tidak ada baris provinsi (3500)", (ind.kode_wilayah != "3500").all())
    dup = ind.duplicated(["kode_wilayah", "tahun", "indikator"]).sum()
    check(g, "key kode_wilayah+tahun+indikator unik", dup == 0, f"{dup} duplikat")
    nreg = ind.groupby(["indikator", "tahun"]).kode_wilayah.nunique()
    check(g, "38 wilayah di setiap indikator×tahun", (nreg == 38).all(),
          f"min {nreg.min()}, max {nreg.max()}")
    na = ind[ind.nilai.isna()]
    check(g, "nilai kosong selalu ber-flag 'tidak_tersedia'",
          na.flag.fillna("").str.contains("tidak_tersedia").all(),
          f"{len(na)} sel kosong: " + ", ".join(sorted({f'{r.indikator} {r.tahun}' for r in na.itertuples()})))
    rng = {"persen_penduduk_miskin": (0, 100), "tpt": (0, 100), "ipm": (40, 100), "rls": (2, 15),
           "hls": (8, 18), "uhh": (60, 80), "gini_ratio": (0, 1),
           "pengeluaran_per_kapita_disesuaikan": (4000, 25000), "jumlah_penduduk": (100_000, 3_500_000),
           "kepadatan_penduduk": (200, 10_000), "pdrb_per_kapita_adhb": (5000, 700_000),
           "laju_pertumbuhan_pdrb_adhk": (-25, 25), "jumlah_penduduk_miskin": (1, 400)}
    bad = []
    for k, (lo, hi) in rng.items():
        s = ind[ind.indikator == k].nilai.dropna()
        if not s.between(lo, hi).all():
            bad.append(f"{k} [{s.min()}, {s.max()}]")
    check(g, "rentang nilai masuk akal (cek konversi ribuan)", not bad, "; ".join(bad) or "semua dalam rentang")
    yrs = ind.groupby("indikator").tahun.agg(["min", "max"])
    warn(g, "cakupan tahun per indikator",
         "; ".join(f"{i}: {r['min']}–{r['max']}" for i, r in yrs.iterrows()))
    flags = ind.flag.fillna("").str.split(";").explode()
    warn(g, "jumlah flag", ", ".join(f"{k}={v}" for k, v in flags[flags != ""].value_counts().items()))

    # C. Provinsi
    g = "C. Indikator provinsi"
    check(g, "key unik", not prov.duplicated(["kode_wilayah", "tahun", "indikator"]).any())
    p = ind[ind.indikator == "jumlah_penduduk"].groupby("tahun").nilai.sum()
    pj = prov[prov.indikator == "jumlah_penduduk"].set_index("tahun").nilai
    d = (p - pj)
    check(g, "Σ penduduk 38 kab/kota = baris Jatim (toleransi 10.000 jiwa)", (d.abs() <= 10_000).all(),
          "selisih maks " + f"{d.abs().max():,.0f} ({d.abs().idxmax()})")
    pm = ind[ind.indikator == "jumlah_penduduk_miskin"].groupby("tahun").nilai.sum()
    pmj = prov[prov.indikator == "jumlah_penduduk_miskin"].set_index("tahun").nilai
    dd = (pm - pmj).round(2)
    warn(g, "Σ jumlah miskin vs baris Jatim",
         "tidak diagregasi; selisih 2010–2014: " + ", ".join(f"{t}: {v:+}" for t, v in dd.items() if abs(v) > 1))

    # D. Konsistensi IPM (validasi UHH, HLS, RLS, PPK)
    g = "D. Konsistensi IPM"
    w = ind.pivot_table(index=["kode_wilayah", "tahun"], columns="indikator", values="nilai")
    Ih = (w.uhh - 20) / 65
    Ie = (w.hls / 18 + w.rls / 15) / 2
    Ix = (np.log(w.pengeluaran_per_kapita_disesuaikan) - np.log(1007.436)) / (np.log(26572.352) - np.log(1007.436))
    dif = (100 * (Ih * Ie * Ix) ** (1 / 3) - w.ipm).abs()
    check(g, "IPM dapat direkonstruksi dari komponennya (≤0,02)", dif.max() <= 0.02,
          f"selisih maks {dif.max():.3f} (rata-rata {dif.mean():.4f})")

    # E. PDRB
    g = "E. PDRB lapangan usaha"
    check(g, "jumlah baris 38×17×5 = 3230", len(pdrb) == 3230, str(len(pdrb)))
    check(g, "key kode_wilayah+tahun+kode_LU unik",
          not pdrb.duplicated(["kode_wilayah", "tahun", "kode_lapangan_usaha"]).any())
    check(g, "seluruh 17 kategori termapping ke sektor", pdrb.sektor.notna().all()
          and set(pdrb.kode_lapangan_usaha) == set(KODE_LU))
    check(g, "38 kode valid", set(pdrb.kode_wilayah) == expected)
    s = pdrb.groupby(["kode_wilayah", "tahun"])[["pdrb_adhb", "pdrb_adhk"]].sum().join(
        tot.set_index(["kode_wilayah", "tahun"]))
    eb = (s.pdrb_adhb - s.pdrb_adhb_total_resmi).abs()
    ek = (s.pdrb_adhk - s.pdrb_adhk_total_resmi).abs()
    check(g, "Σ17 LU ADHB = total resmi (≤0,5 miliar)", eb.max() <= 0.5, f"maks {eb.max():.2f} ({eb.idxmax()[0]}, {int(eb.idxmax()[1])})")
    check(g, "Σ17 LU ADHK = total resmi (≤0,5 miliar)", ek.max() <= 0.5, f"maks {ek.max():.2f} ({ek.idxmax()[0]}, {int(ek.idxmax()[1])})")
    defl = s.pdrb_adhb / s.pdrb_adhk
    check(g, "deflator ADHB/ADHK dalam 0,9–2,5", defl.between(0.9, 2.5).all(),
          f"min {defl.min():.3f} ({defl.idxmin()[0]}, {int(defl.idxmin()[1])}), max {defl.max():.3f} ({defl.idxmax()[0]}, {int(defl.idxmax()[1])})")
    same = pdrb.groupby("kode_wilayah").apply(lambda x: (x.pdrb_adhb == x.pdrb_adhk).all())
    check(g, "tidak ada wilayah dengan ADHK = salinan ADHB", not same.any(), ", ".join(same[same].index))
    # silang dengan lampiran indeks implisit per LU
    pm_ = pd.read_csv(REPORTS / "pdrb_pdf_page_map.csv", dtype={"kode_wilayah": str})
    lu = pd.read_csv(INTERIM / "pdf_lu_tables.csv")
    imp = pm_[pm_.jenis == "IMPLISIT"]
    errs = []
    for r in imp.itertuples():
        I = lu[lu.halaman_pdf == r.halaman_pdf].pivot(index="kode_lu", columns="tahun", values="nilai")
        x = pdrb[pdrb.kode_wilayah == r.kode_wilayah].pivot(index="kode_lapangan_usaha", columns="tahun",
                                                             values=["pdrb_adhb", "pdrb_adhk"])
        rat = (x.pdrb_adhb / x.pdrb_adhk * 100)
        e = ((rat - I.reindex(rat.index)).abs() / I.reindex(rat.index))
        e = e[(x.pdrb_adhk > 5)]  # LU sangat kecil sensitif pembulatan
        errs.append((r.kode_wilayah, float(np.nanmax(e.values))))
    worst = max(errs, key=lambda t: t[1])
    check(g, f"ADHB/ADHK per LU = indeks implisit lampiran ({len(errs)} wilayah, ≤1%)", worst[1] <= 0.01,
          f"deviasi relatif maks {worst[1]*100:.2f}% ({worst[0]})")
    sh = pm_[pm_.jenis == "IMPLISIT_BERGESER_1_KOLOM"].kode_wilayah.tolist()
    warn(g, "lampiran indeks implisit bergeser kolom di publikasi (tidak dipakai)", ", ".join(sh))
    warn(g, "status data", "2023 = sementara, 2024 = sangat sementara (kolom status_data)")

    # F. GeoJSON
    g = "F. GeoJSON"
    gj = json.load(open(GEO / "jatim_kabkota.geojson", encoding="utf-8"))
    codes = [f["properties"]["kode_wilayah"] for f in gj["features"]]
    check(g, "38 fitur, kode unik", len(codes) == 38 and len(set(codes)) == 38)
    check(g, "kode GeoJSON = tabel kode wilayah", set(codes) == set(kode.kode_wilayah))
    try:
        from shapely.geometry import shape
        inval = [f["properties"]["nama_wilayah"] for f in gj["features"] if not shape(f["geometry"]).is_valid]
        check(g, "geometri valid (shapely)", not inval, ", ".join(inval))
    except ImportError:
        warn(g, "geometri valid", "shapely tidak terpasang; dilewati")
    j = ind[(ind.indikator == "persen_penduduk_miskin") & (ind.tahun == 2025)]
    check(g, "join indikator 2025 ↔ GeoJSON 38/38", set(j.kode_wilayah) == set(codes))

    # G. Kesiapan multivariat
    g = "G. Kesiapan PCA"
    vars_ = ["persen_penduduk_miskin", "pengeluaran_per_kapita_disesuaikan", "rls", "hls", "uhh", "tpt",
             "pdrb_per_kapita_adhb", "laju_pertumbuhan_pdrb_adhk", "kepadatan_penduduk", "gini_ratio"]
    for th in (2024, 2025):
        m = ind[(ind.tahun == th) & ind.indikator.isin(vars_)].pivot(index="kode_wilayah", columns="indikator",
                                                                     values="nilai")
        check(g, f"tahun {th}: ≥8 variabel × ≥34 unit, tanpa NA",
              m.shape[1] >= 8 and m.shape[0] >= 34 and not m.isna().any().any(), f"{m.shape[0]} unit × {m.shape[1]} variabel")

    # H. Sheet duplikat PPK (D02 vs D05)
    g = "H. Catatan sumber"
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    a, b = wb["per kapita"], wb["pengeluaran per kapita jatim"]
    diffs = [(a.cell(r, 1).value, int(a.cell(4, c).value)) for r in range(5, 44) for c in range(2, 18)
             if a.cell(r, c).value != b.cell(r, c).value]
    warn(g, "sheet D02 vs D05 (dipakai D05)", f"{len(diffs)} sel beda: {diffs}")

    # tulis laporan
    df = pd.DataFrame(results, columns=["kelompok", "pemeriksaan", "status", "detail"])
    nfail = (df.status == "FAIL").sum()
    with open(REPORTS / "validation_report.md", "w", encoding="utf-8") as fh:
        fh.write(f"# Laporan Validasi Data\n\nPASS: {(df.status == 'PASS').sum()} | "
                 f"FAIL: {nfail} | CATATAN: {(df.status == 'CATATAN').sum()}\n\n")
        for grp, sub in df.groupby("kelompok", sort=False):
            fh.write(f"## {grp}\n\n| Pemeriksaan | Status | Detail |\n|---|---|---|\n")
            for r in sub.itertuples():
                fh.write(f"| {r.pemeriksaan} | {r.status} | {str(r.detail).replace('|', '/')} |\n")
            fh.write("\n")
    print(df.to_string(index=False, max_colwidth=90))
    print(f"\n[05] PASS {(df.status == 'PASS').sum()} | FAIL {nfail}")
    sys.exit(1 if nfail else 0)


if __name__ == "__main__":
    main()
