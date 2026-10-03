"""STEP 1b — Tabel kode wilayah (D21).

Sumber: sheet 'KODE WILAYAH' (Wilkerstat BPS). Kode disimpan sebagai float
di Excel (3501.0) dan nama tanpa awalan Kab/Kota (7 nama kembar), sehingga:
  - kode -> string 4 digit
  - jenis_wilayah: kode < 3571 = Kabupaten, >= 3571 = Kota
  - nama_wilayah  : 'Kabupaten X' / 'Kota X'
  - nama_tabel    : nama seperti tertulis di sheet indikator BPS
                    ('Pacitan' untuk kabupaten, 'Kota Kediri' untuk kota)
"""
import openpyxl
import pandas as pd
from common import XLSX, PROCESSED, ensure_dirs


def main():
    ensure_dirs()
    ws = openpyxl.load_workbook(XLSX, data_only=True)["KODE WILAYAH"]
    rows = []
    for r in range(4, ws.max_row + 1):
        nama, kode = ws.cell(r, 2).value, ws.cell(r, 3).value
        if nama is None or kode is None:
            continue
        kode = f"{int(kode):04d}"
        nama = str(nama).strip()
        jenis = "Kabupaten" if int(kode) < 3571 else "Kota"
        rows.append({
            "kode_wilayah": kode,
            "nama_wilayah": f"{jenis} {nama}",
            "jenis_wilayah": jenis,
            "nama_singkat": nama,
            "nama_tabel": nama if jenis == "Kabupaten" else f"Kota {nama}",
            "nama_pdf": f"Kab. {nama}" if jenis == "Kabupaten" else f"Kota {nama}",
        })
    df = pd.DataFrame(rows)
    df.to_csv(PROCESSED / "kode_wilayah.csv", index=False)
    print(f"[01] kode_wilayah: {len(df)} baris")


if __name__ == "__main__":
    main()
